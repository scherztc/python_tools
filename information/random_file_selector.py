import os
import random
import PyPDF2
import docx
import openpyxl
from PIL import Image
import pygame
import cv2

def get_random_file(folder_path):
    files = [f for f in os.listdir(folder_path) if os.path.isfile(os.path.join(folder_path, f))]
    if not files:
        raise Exception("No files found in the specified folder.")
    random_file = random.choice(files)
    return os.path.join(folder_path, random_file)

def open_file(file_path):
    extension = os.path.splitext(file_path)[1].lower()
    content = ""

    if extension == ".txt":
        with open(file_path, 'r') as file:
            content = file.read()
    elif extension == ".pdf":
        with open(file_path, 'rb') as file:
            reader = PyPDF2.PdfFileReader(file)
            for page_num in range(reader.numPages):
                page = reader.getPage(page_num)
                content += page.extract_text() + "\n"
    elif extension == ".docx":
        doc = docx.Document(file_path)
        for para in doc.paragraphs:
            content += para.text + "\n"
    elif extension == ".xlsx":
        workbook = openpyxl.load_workbook(file_path)
        sheet = workbook.active
        for row in sheet.iter_rows(values_only=True):
            content += '\t'.join([str(cell) for cell in row]) + '\n'
    elif extension in [".jpg", ".jpeg", ".png", ".bmp", ".gif"]:
        img = Image.open(file_path)
        img.show()
        content = "Image opened successfully."
    elif extension in [".mp3", ".wav", ".ogg"]:
        pygame.mixer.init()
        pygame.mixer.music.load(file_path)
        pygame.mixer.music.play()
        content = "Playing audio file."
    elif extension in [".mp4", ".avi", ".mov"]:
        cap = cv2.VideoCapture(file_path)
        while cap.isOpened():
            ret, frame = cap.read()
            if not ret:
                break
            cv2.imshow('Video', frame)
            if cv2.waitKey(25) & 0xFF == ord('q'):
                break
        cap.release()
        cv2.destroyAllWindows()
        content = "Video played successfully."
    else:
        content = "Unsupported file type."

    return content

def main():
    folder_path = input("Enter the path to the folder: ")
    try:
        random_file_path = get_random_file(folder_path)
        print(f"Randomly selected file: {random_file_path}")
        content = open_file(random_file_path)
        print("File content:\n", content)
    except Exception as e:
        print(f"An error occurred: {e}")

if __name__ == "__main__":
    main()

